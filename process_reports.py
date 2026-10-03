"""Bluebonnet Blend Co. weekly sales pipeline (Parts 1-4).

One run does four things, in order:

1. Consolidate -- scan a folder (every subfolder and every .zip inside it, at
   any depth), turn each weekly report into one row and validate its values.
2. Store -- upsert those rows into Sales_Master.xlsx. History is never
   overwritten; a report for a store and week already stored replaces that row.
3. Analyse -- momentum, 26-week highs/lows, weekly swings and rankings, based on
   each store's most recent entry in the master.
4. Report -- write weekly_ops_brief.txt.

Usage (run once per weekly drop, in the order the drops arrive):
    uv run process_reports.py Input_Files
    uv run process_reports.py Next_Week_Drop

Design decisions worth knowing before reading the code:

* The filename is the source of truth for Store Code and Week Ending.
* A file that cannot produce a row (not a report, store not in the directory,
  unreadable) is logged and skipped. Nothing in a single file can stop the run.
* A value that cannot be true is *blanked* in the master and its original value
  is written to the Data Quality Flag column. Anything that reads the master
  (analysis, dashboard, the agent) therefore never sees a flagged value, while
  the flag text keeps the row auditable. The rest of the row stays usable.
* Later runs win: re-running an old drop after a correction would put the old
  numbers back. Run drops in arrival order.
* 26-week highs/lows for stores with less history: see ``store_snapshot``.
"""

import argparse
import io
import logging
import math
import os
import re
import sys
import zipfile
from collections import Counter
from collections.abc import Callable, Iterator
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

log = logging.getLogger("process_reports")

# --- Output layout -----------------------------------------------------------

METRICS = ("Net Sales", "Transactions", "Avg Ticket")
PERIODS = ("7D", "28D")


def column_name(period: str, metric: str, pct: bool = False) -> str:
    """Output column for one measure, e.g. '28D Net Sales % Chg'."""
    return f"{period} {metric}{' % Chg' if pct else ''}"


# Where each measure lives on the Summary tab (see Summary_Tab_Screenshot.png).
SUMMARY_SHEET = "Summary"
METRIC_ROWS = {"Net Sales": 11, "Transactions": 12, "Avg Ticket": 13}
MEASURE_LETTERS = {("7D", False): "C", ("7D", True): "E", ("28D", False): "G", ("28D", True): "I"}
EXCEL_CELLS = {
    column_name(period, metric, pct): f"{letter}{row}"
    for metric, row in METRIC_ROWS.items()
    for (period, pct), letter in MEASURE_LETTERS.items()
}

METRIC_COLUMNS = list(EXCEL_CELLS)
OUTPUT_COLUMNS = ["Store Code", "Week Ending", *METRIC_COLUMNS, "Data Quality Flag"]

# The long-layout CSV names the same things differently.
CSV_REQUIRED = {"metric", "period", "this_year", "pct_change_vs_ly"}
CSV_METRICS = {metric.casefold(): metric for metric in METRICS}
CSV_PERIODS = {"7 day": "7D", "28 day": "28D"}

# --- Filename formats --------------------------------------------------------

# Format 1: DAL01-20260830.xlsx
CODE_FILE = re.compile(r"^(?P<code>[A-Za-z0-9]+)-(?P<date>\d{8})\.xlsx$", re.IGNORECASE)
# Format 2: Weekly Sales_Austin Mueller-20260830.xlsx
NAME_FILE = re.compile(r"^Weekly Sales_(?P<name>.+)-(?P<date>\d{8})\.xlsx$", re.IGNORECASE)
# Format 3: 2026-08-30_sat01_export.csv
EXPORT_FILE = re.compile(r"^(?P<date>\d{4}-\d{2}-\d{2})_(?P<code>[A-Za-z0-9]+)_export\.csv$", re.IGNORECASE)

# --- Why a file was skipped --------------------------------------------------

NOT_A_REPORT = "not a report"
UNMAPPED_STORE = "unmapped store"
UNREADABLE = "unreadable"
SUPERSEDED = "superseded"

# Avg Ticket is rounded to cents, so Net Sales can drift from
# Transactions x Avg Ticket by up to half a cent per transaction.
TICKET_TOLERANCE = 0.01


class SkipFile(Exception):
    """A file cannot contribute a row. Carries the category for the run summary."""

    def __init__(self, category: str, reason: str):
        super().__init__(reason)
        self.category = category


@dataclass
class SourceFile:
    """A file on disk or a member of a (possibly nested) zip archive."""

    path: str  # relative to the scanned folder; zip members read as outer.zip/member
    read: Callable[[], bytes]
    modified: datetime

    @property
    def name(self) -> str:
        return self.path.rsplit("/", 1)[-1]

    @property
    def is_correction(self) -> bool:
        return "corrections" in (part.casefold() for part in self.path.split("/")[:-1])


@dataclass
class Report:
    """One consolidated row plus the provenance needed to resolve duplicates."""

    store_code: str
    week_ending: date
    values: dict[str, float | None]
    flags: list[str]
    source: SourceFile

    @property
    def rank(self) -> tuple[bool, datetime]:
        """Higher wins when two files in one run report the same store and week."""
        return (self.source.is_correction, self.source.modified)

    def to_row(self) -> dict[str, object]:
        return {
            "Store Code": self.store_code,
            "Week Ending": self.week_ending,
            **self.values,
            "Data Quality Flag": "; ".join(self.flags),
        }


@dataclass
class ScanResult:
    """Everything one run learned: the rows, and what was left out and why."""

    reports: dict[tuple[str, date], Report] = field(default_factory=dict)
    skipped: list[tuple[str, str, str]] = field(default_factory=list)  # (path, category, reason)
    files_seen: int = 0

    def skip(self, path: str, category: str, reason: str) -> None:
        self.skipped.append((path, category, reason))
        log.warning("Skipped %s [%s]: %s", path, category, reason)

    def to_frame(self) -> pd.DataFrame:
        rows = [report.to_row() for report in self.reports.values()]
        frame = pd.DataFrame(rows, columns=OUTPUT_COLUMNS)
        frame["Week Ending"] = pd.to_datetime(frame["Week Ending"])
        return frame.sort_values(["Store Code", "Week Ending"], ignore_index=True)


# --- Finding files -----------------------------------------------------------


def iter_sources(folder: Path, result: ScanResult) -> Iterator[SourceFile]:
    """Yield every file under ``folder``, looking inside zip archives at any depth."""
    for path in sorted(folder.rglob("*")):
        if path.is_file():
            modified = datetime.fromtimestamp(path.stat().st_mtime)
            source = SourceFile(path.relative_to(folder).as_posix(), path.read_bytes, modified)
            yield from expand(source, result)


def expand(source: SourceFile, result: ScanResult) -> Iterator[SourceFile]:
    """Yield the file itself, or the members of the archive if it is a .zip."""
    if not source.name.lower().endswith(".zip"):
        yield source
        return
    try:
        # Archives are opened in memory so a zip inside a zip needs no temp files.
        archive = zipfile.ZipFile(io.BytesIO(source.read()))
    except Exception as exc:  # corrupt, encrypted or unreadable archive
        result.files_seen += 1
        result.skip(source.path, UNREADABLE, f"cannot open archive: {exc}")
        return
    with archive:
        for info in archive.infolist():
            if info.is_dir():
                continue
            member = SourceFile(
                f"{source.path}/{info.filename.replace('\\', '/')}",
                lambda info=info: archive.read(info),
                datetime(*info.date_time),
            )
            yield from expand(member, result)


# --- Step 1: parse the filename ----------------------------------------------


def normalise_name(name: str) -> str:
    return " ".join(name.split()).casefold()


def load_directory(path: Path) -> dict[str, str]:
    """Return {store code: official store name} from Store_Directory.csv."""
    table = pd.read_csv(path, dtype=str).dropna(subset=["store_code", "store_name"])
    return {code.strip().upper(): name.strip() for code, name in zip(table["store_code"], table["store_name"])}


def parse_filename(name: str, directory: dict[str, str]) -> tuple[str, date, Callable[[bytes], dict]]:
    """Return (store code, week ending, reader for this layout) or raise SkipFile."""
    if name.startswith("~$"):
        raise SkipFile(NOT_A_REPORT, "Office lock file left behind by an open workbook")

    if match := NAME_FILE.match(name):
        wanted = normalise_name(match["name"])
        code = next((code for code, official in directory.items() if normalise_name(official) == wanted), None)
        if code is None:
            raise SkipFile(UNMAPPED_STORE, f"store name {match['name']!r} is not in the store directory")
        date_format, reader = "%Y%m%d", read_excel_summary
    elif match := CODE_FILE.match(name):
        code, date_format, reader = match["code"].upper(), "%Y%m%d", read_excel_summary
    elif match := EXPORT_FILE.match(name):
        code, date_format, reader = match["code"].upper(), "%Y-%m-%d", read_long_csv
    else:
        raise SkipFile(NOT_A_REPORT, "filename does not match any report format")

    # The directory is the official store list, so an unknown code is treated
    # exactly like an unknown name.
    if code not in directory:
        raise SkipFile(UNMAPPED_STORE, f"store code {code} is not in the store directory")
    try:
        week_ending = datetime.strptime(match["date"], date_format).date()
    except ValueError:
        raise SkipFile(NOT_A_REPORT, f"{match['date']!r} in the filename is not a real date") from None
    return code, week_ending, reader


# --- Step 2: extract the 12 metrics ------------------------------------------


def read_excel_summary(data: bytes) -> dict[str, object]:
    """Read the 12 cells from the Summary tab of an Excel export."""
    try:
        workbook = load_workbook(io.BytesIO(data), data_only=True)
    except Exception as exc:  # openpyxl raises many different types for a damaged file
        raise SkipFile(UNREADABLE, f"cannot open workbook: {exc}") from exc
    if SUMMARY_SHEET not in workbook.sheetnames:
        raise SkipFile(UNREADABLE, f"workbook has no {SUMMARY_SHEET!r} tab")
    sheet = workbook[SUMMARY_SHEET]
    # Fixed cell addresses are only safe if the table is where we expect it.
    for metric, row in METRIC_ROWS.items():
        label = sheet[f"B{row}"].value
        if str(label).strip().casefold() != metric.casefold():
            raise SkipFile(UNREADABLE, f"unexpected layout: B{row} is {label!r}, expected {metric!r}")
    return {column: sheet[cell].value for column, cell in EXCEL_CELLS.items()}


def read_long_csv(data: bytes) -> dict[str, object]:
    """Read the same 12 values from a long-layout CSV (one row per metric and period)."""
    try:
        table = pd.read_csv(io.BytesIO(data), dtype=str, encoding="utf-8-sig")
    except Exception as exc:  # empty file, wrong encoding, ragged rows
        raise SkipFile(UNREADABLE, f"cannot parse CSV: {exc}") from exc
    table.columns = table.columns.str.strip().str.lower()
    if missing := CSV_REQUIRED - set(table.columns):
        raise SkipFile(UNREADABLE, f"CSV is missing columns {sorted(missing)}")

    raw: dict[str, object] = {}
    for row in table.itertuples(index=False):
        metric = CSV_METRICS.get(str(row.metric).strip().casefold())
        period = CSV_PERIODS.get(str(row.period).strip().casefold())
        if metric is None or period is None:
            continue  # Labor % and anything else we were not asked for
        level = column_name(period, metric)
        if level in raw:
            raise SkipFile(UNREADABLE, f"CSV has more than one {period} {metric} row")
        raw[level] = row.this_year
        raw[column_name(period, metric, pct=True)] = row.pct_change_vs_ly
    if not raw:
        raise SkipFile(UNREADABLE, "CSV has none of the expected metric rows")
    return raw


# --- Step 3: validate --------------------------------------------------------


def to_number(raw: object) -> float | None:
    """Coerce a cell or CSV value to a float; None if it is blank or not a finite number."""
    if isinstance(raw, bool):
        return None
    if isinstance(raw, str):
        try:
            raw = float(raw.replace("$", "").replace(",", "").replace("%", "").strip())
        except ValueError:
            return None
    if not isinstance(raw, (int, float)) or not math.isfinite(raw):
        return None
    return float(raw)


def validate(raw: dict[str, object]) -> tuple[dict[str, float | None], list[str]]:
    """Return (values with impossible ones blanked, one flag per blanked value).

    A value is flagged only when it *cannot* be true. Unusual-but-possible
    numbers (a store up 40% on last year) are left alone: that is news, not noise.
    """
    values = {column: to_number(raw.get(column)) for column in METRIC_COLUMNS}
    problems: dict[str, str] = {}

    def flag(column: str, reason: str) -> None:
        problems.setdefault(column, reason)  # keep the first, most specific reason

    for period in PERIODS:
        for metric in METRICS:
            level, pct = column_name(period, metric), column_name(period, metric, pct=True)
            if values[level] is None:
                flag(level, "missing or not a number")
            elif values[level] < 0:
                flag(level, "negative value")
            elif metric == "Transactions" and not values[level].is_integer():
                flag(level, "not a whole number")

            if values[pct] is None:
                flag(pct, "missing or not a number")
            elif values[pct] < -100:
                # Sales, transactions and tickets cannot go below zero, so they
                # cannot fall by more than 100% against last year.
                flag(pct, "below -100%")
            elif level in problems:
                flag(pct, "change is based on an invalid value")

        # Net Sales = Transactions x Avg Ticket. If the three disagree we cannot
        # tell which one is wrong, so none of them can be trusted.
        sales, count, ticket = (column_name(period, metric) for metric in METRICS)
        if not problems.keys() & {sales, count, ticket}:
            drift = abs(values[sales] - values[count] * values[ticket])
            if drift > TICKET_TOLERANCE * max(values[count], 1):
                for column in (sales, count, ticket):
                    flag(column, "Net Sales does not equal Transactions x Avg Ticket")

    # The 28-day window contains the 7-day window, so its total cannot be smaller.
    for metric in ("Net Sales", "Transactions"):
        short, long = column_name("7D", metric), column_name("28D", metric)
        if short not in problems and long not in problems and values[long] < values[short]:
            flag(short, "larger than the 28D total that contains it")
            flag(long, "smaller than the 7D total it contains")

    flags = [f"{column}: {problems[column]} ({raw.get(column)!r})" for column in METRIC_COLUMNS if column in problems]
    clean = {column: None if column in problems else value for column, value in values.items()}
    return clean, flags


# --- Putting it together -----------------------------------------------------


def read_report(source: SourceFile, directory: dict[str, str]) -> Report:
    """Turn one file into a Report, or raise SkipFile explaining why not."""
    store_code, week_ending, reader = parse_filename(source.name, directory)
    try:
        data = source.read()
    except Exception as exc:  # disk error, or a damaged/encrypted zip member
        raise SkipFile(UNREADABLE, f"cannot read file: {exc}") from exc
    values, flags = validate(reader(data))
    return Report(store_code, week_ending, values, flags, source)


def process_folder(folder: Path, directory: dict[str, str]) -> ScanResult:
    """Scan ``folder`` and consolidate every report in it into a ScanResult."""
    result = ScanResult()
    for source in iter_sources(folder, result):
        result.files_seen += 1
        try:
            report = read_report(source, directory)
        except SkipFile as skip:
            result.skip(source.path, skip.category, str(skip))
            continue

        for flag_text in report.flags:
            log.warning("Flagged %s %s (%s): %s", report.store_code, report.week_ending, source.path, flag_text)
        log.debug("Read %s -> %s %s", source.path, report.store_code, report.week_ending)

        # One row per store and week. If a run contains the same report twice,
        # a file in a Corrections folder wins, then the most recently modified.
        key = (report.store_code, report.week_ending)
        previous = result.reports.get(key)
        if previous is not None and previous.rank > report.rank:
            report, previous = previous, report
        result.reports[key] = report
        if previous is not None:
            result.skip(previous.source.path, SUPERSEDED, f"replaced by {report.source.path}")
    return result


# --- Part 2: the master file -------------------------------------------------

KEY = ["Store Code", "Week Ending"]
MASTER_SHEET = "Sales Master"


@dataclass
class StoreResult:
    """The master after an upsert, and what the upsert changed."""

    master: pd.DataFrame
    added: list[tuple[str, pd.Timestamp]]
    revised: list[tuple[str, pd.Timestamp]]  # already stored; the new report changed it
    unchanged: int  # already stored with identical values (e.g. a drop processed twice)


def load_master(path: Path) -> pd.DataFrame:
    """Read Sales_Master.xlsx, or start an empty master if there is none yet."""
    if not path.exists():
        return pd.DataFrame(columns=OUTPUT_COLUMNS).astype({"Week Ending": "datetime64[ns]"})
    master = pd.read_excel(path, sheet_name=0, dtype={"Store Code": str, "Data Quality Flag": str})
    if missing := [column for column in OUTPUT_COLUMNS if column not in master.columns]:
        raise ValueError(f"{path.name} is missing columns {missing}")
    master = master[OUTPUT_COLUMNS].drop_duplicates(KEY, keep="last")
    master["Week Ending"] = pd.to_datetime(master["Week Ending"])
    master[METRIC_COLUMNS] = master[METRIC_COLUMNS].apply(pd.to_numeric, errors="coerce")
    master["Data Quality Flag"] = master["Data Quality Flag"].fillna("")
    return master


def same_report(a: pd.Series, b: pd.Series) -> bool:
    """True when two rows for one store-week carry identical values (blank == blank)."""
    values_match = all(
        (pd.isna(x) and pd.isna(y)) or x == y for x, y in zip(a[METRIC_COLUMNS], b[METRIC_COLUMNS])
    )
    return values_match and a["Data Quality Flag"] == b["Data Quality Flag"]


def upsert(master: pd.DataFrame, new: pd.DataFrame) -> StoreResult:
    """Append new store-weeks and replace stored ones that arrived again.

    Rows of the master that this run did not touch are carried over as they are,
    so history is never lost. A resubmitted report replaces the stored row
    rather than adding a second one.
    """
    stored, incoming = master.set_index(KEY), new.set_index(KEY)
    added, revised = [], []
    for key, row in incoming.iterrows():
        if key not in stored.index:
            added.append(key)
        elif not same_report(row, stored.loc[key]):
            revised.append(key)
    kept = stored.drop(index=incoming.index, errors="ignore")
    combined = pd.concat([kept, incoming]) if len(kept) else incoming
    combined = combined.sort_index().reset_index()[OUTPUT_COLUMNS]
    return StoreResult(combined, added, revised, len(incoming) - len(added) - len(revised))


def excel_format(column: str) -> str | None:
    if column == "Week Ending":
        return "yyyy-mm-dd"
    if column.endswith("% Chg"):
        return "0.0"
    for metric, number_format in (("Net Sales", "#,##0.00"), ("Transactions", "#,##0"), ("Avg Ticket", "0.00")):
        if metric in column:
            return number_format
    return None


def save_master(master: pd.DataFrame, path: Path) -> None:
    """Write the master to a temp file and swap it in, so a failed write cannot corrupt history."""
    temp = path.with_name(f"{path.stem}.tmp{path.suffix}")
    try:
        with pd.ExcelWriter(temp, engine="openpyxl") as writer:
            master.to_excel(writer, sheet_name=MASTER_SHEET, index=False)
            sheet = writer.sheets[MASTER_SHEET]
            sheet.freeze_panes = "C2"
            sheet.auto_filter.ref = sheet.dimensions
            for index, column in enumerate(master.columns, start=1):
                letter = get_column_letter(index)
                sheet.column_dimensions[letter].width = 70 if column == "Data Quality Flag" else max(12, len(column) + 2)
                if number_format := excel_format(column):
                    for (cell,) in sheet.iter_rows(min_row=2, min_col=index, max_col=index):
                        cell.number_format = number_format
        os.replace(temp, path)
    finally:
        temp.unlink(missing_ok=True)


# --- Part 3: analysis --------------------------------------------------------

SWING_ALERT_POINTS = 5.0  # a week-over-week move in 7D Net Sales % Chg bigger than this is urgent
TREND_WEEKS = 26
MIN_TREND_WEEKS = 4  # with less history than this, no high or low is called at all

SNAPSHOT_COLUMNS = [
    "Store Code", "Week Ending", "7D Net Sales", "7D Net Sales % Chg", "Previous 7D Net Sales % Chg",
    "Swing (pts)", "28D Net Sales", "28D Net Sales % Chg", "28D Transactions % Chg",
    "28D Avg Ticket % Chg", "28D Driver", "Trend", "Trend Weeks", "Previous High", "Previous Low",
]


def sales_driver(transactions_pct: float, ticket_pct: float) -> str | None:
    """Which half of Net Sales = Transactions x Avg Ticket moved Net Sales more."""
    if pd.isna(transactions_pct) or pd.isna(ticket_pct):
        return None
    return "traffic" if abs(transactions_pct) >= abs(ticket_pct) else "average ticket"


def store_snapshot(master: pd.DataFrame) -> pd.DataFrame:
    """One row per store, built from that store's most recent entry.

    * Swing: this week's 7D Net Sales % Chg minus the store's value exactly one
      week earlier. Blank when that week is missing or its value was flagged.
    * Trend: whether 28D Net Sales is at its highest or lowest point of the
      trailing 26 weeks (this week included).

    Stores with less than 26 weeks of history (ELP01 opened in July) are judged
    over the weeks they do have, and ``Trend Weeks`` records how many -- so a
    9-week low is never presented as a 26-week low. With fewer than
    MIN_TREND_WEEKS weeks no high or low is called. Missing or flagged weeks
    simply drop out of the window.
    """
    rows = []
    for code, history in master.groupby("Store Code"):
        history = history.sort_values("Week Ending")
        latest = history.iloc[-1]
        week = latest["Week Ending"]

        previous = history.loc[history["Week Ending"] == week - pd.Timedelta(weeks=1), "7D Net Sales % Chg"]
        previous_pct = previous.iloc[0] if len(previous) else math.nan
        # Rounded so float noise cannot turn a 5.0-point move into an alert.
        swing = round(latest["7D Net Sales % Chg"] - previous_pct, 1)

        window = history[history["Week Ending"] > week - pd.Timedelta(weeks=TREND_WEEKS)]
        earlier = window.loc[window["Week Ending"] < week, "28D Net Sales"].dropna()
        trend_weeks = min(TREND_WEEKS, (week - history["Week Ending"].min()).days // 7 + 1)
        current, trend = latest["28D Net Sales"], None
        if pd.notna(current) and trend_weeks >= MIN_TREND_WEEKS and len(earlier):
            if current >= earlier.max():
                trend = "HIGH"
            elif current <= earlier.min():
                trend = "LOW"

        rows.append([
            code, week, latest["7D Net Sales"], latest["7D Net Sales % Chg"], previous_pct, swing,
            current, latest["28D Net Sales % Chg"], latest["28D Transactions % Chg"],
            latest["28D Avg Ticket % Chg"],
            sales_driver(latest["28D Transactions % Chg"], latest["28D Avg Ticket % Chg"]),
            trend, trend_weeks,
            earlier.max() if len(earlier) else math.nan, earlier.min() if len(earlier) else math.nan,
        ])
    return pd.DataFrame(rows, columns=SNAPSHOT_COLUMNS)


@dataclass
class Analysis:
    """Part 3 results. Shared by the brief, the dashboard and the agent so all three agree."""

    week: pd.Timestamp | None  # the most recent week in the master
    snapshot: pd.DataFrame  # one row per store, see store_snapshot
    flagged: pd.DataFrame  # master rows that carry a Data Quality Flag
    missing_weeks: dict[str, list[date]]  # gaps inside each store's run of weeks

    @property
    def comparable(self) -> pd.DataFrame:
        return self.snapshot.dropna(subset=["Swing (pts)"])

    @property
    def improved(self) -> pd.DataFrame:
        """1. Momentum: stores whose 7D Net Sales % Chg rose on the previous week."""
        return self.comparable[self.comparable["Swing (pts)"] > 0].sort_values("Swing (pts)", ascending=False)

    @property
    def trend_alerts(self) -> pd.DataFrame:
        """2. Historical trends: stores at a high or low of their trailing window."""
        return self.snapshot.dropna(subset=["Trend"]).sort_values(["Trend", "Store Code"])

    @property
    def swing_alerts(self) -> pd.DataFrame:
        """3. Urgent: week-over-week swing beyond +/- SWING_ALERT_POINTS."""
        swings = self.comparable
        return swings[swings["Swing (pts)"].abs() > SWING_ALERT_POINTS].sort_values("Swing (pts)")

    def ranked(self) -> pd.DataFrame:
        """4. Rankings: best to worst 28D Net Sales % Chg."""
        return self.snapshot.dropna(subset=["28D Net Sales % Chg"]).sort_values("28D Net Sales % Chg", ascending=False)

    @property
    def stale(self) -> pd.DataFrame:
        """Stores whose latest entry is older than the latest week (they did not report)."""
        return self.snapshot[self.snapshot["Week Ending"] < self.week]


def analyse(master: pd.DataFrame) -> Analysis:
    """Run Part 3 on the master (as loaded by load_master)."""
    missing: dict[str, list[date]] = {}
    for code, history in master.groupby("Store Code"):
        expected = pd.date_range(history["Week Ending"].min(), history["Week Ending"].max(), freq="7D")
        if len(gaps := expected.difference(history["Week Ending"])):
            missing[code] = [gap.date() for gap in gaps]
    return Analysis(
        week=master["Week Ending"].max() if len(master) else None,
        snapshot=store_snapshot(master),
        flagged=master[master["Data Quality Flag"] != ""],
        missing_weeks=missing,
    )


# --- Part 4: the weekly brief ------------------------------------------------


def money(value: float) -> str:
    return f"${value:,.0f}"


def render_brief(analysis: Analysis, names: dict[str, str], scan: ScanResult, stored: StoreResult) -> str:
    """Turn the analysis into a plain-text brief a manager can read in two minutes."""
    rule = "=" * 60
    lines = [rule, "  BLUEBONNET BLEND CO. | WEEKLY OPS BRIEF"]
    if analysis.week is None:
        return "\n".join([*lines, "  No data in the master yet.", rule, ""])
    lines += [f"  Week Ending: {analysis.week.date()}", rule]

    def section(title: str) -> None:
        lines.extend(["", f"--- {title} ---"])

    def store(code: str) -> str:
        return f"{code} {names.get(code, '')}".rstrip()

    snapshot = analysis.snapshot
    section("MOMENTUM (7-Day Net Sales vs LY)")
    lines.append(f"  {len(analysis.improved)} of {len(snapshot)} stores improved on last week.")
    comparable = analysis.comparable.sort_values("Swing (pts)", ascending=False)
    for label, rows in (("Up:  ", comparable[comparable["Swing (pts)"] > 0]), ("Down:", comparable[comparable["Swing (pts)"] <= 0])):
        if len(rows):
            lines.append(f"  {label} " + ", ".join(f"{r['Store Code']} ({r['Swing (pts)']:+.1f})" for _, r in rows.iterrows()))
    if not_comparable := sorted(set(snapshot["Store Code"]) - set(comparable["Store Code"])):
        lines.append(f"  Not comparable (no valid figure for the week before): {', '.join(not_comparable)}")

    section("26-WEEK TREND ALERTS (28-Day Net Sales)")
    for _, row in analysis.trend_alerts.iterrows():
        kind = "high" if row["Trend"] == "HIGH" else "low"
        line = f"  {row['Trend'] + ':':<5} {store(row['Store Code'])} at {money(row['28D Net Sales'])} (previous {kind} {money(row[f'Previous {kind.title()}'])})"
        if row["Trend Weeks"] < TREND_WEEKS:
            line += f" -- {row['Trend Weeks']}-week {kind}, limited history"
        lines.append(line)
    if analysis.trend_alerts.empty:
        lines.append("  No store is at a 26-week high or low.")
    for _, row in snapshot[snapshot["Trend Weeks"] < TREND_WEEKS].iterrows():
        lines.append(f"  Note: {row['Store Code']} has only {row['Trend Weeks']} weeks of history; judged on those weeks.")

    section(f"URGENT: WEEKLY SWING > {SWING_ALERT_POINTS:g} PTS")
    for _, row in analysis.swing_alerts.iterrows():
        lines.append(
            f"  {store(row['Store Code'])}: {row['Swing (pts)']:+.1f} pts "
            f"(this week {row['7D Net Sales % Chg']:+.1f}% | last week {row['Previous 7D Net Sales % Chg']:+.1f}%)"
        )
    if analysis.swing_alerts.empty:
        lines.append(f"  No store moved more than {SWING_ALERT_POINTS:g} points.")

    def ranking(rows: pd.DataFrame) -> None:
        for position, (_, row) in enumerate(rows.iterrows(), start=1):
            why = ""
            if row["28D Driver"]:
                why = f"  driven by {row['28D Driver']} (traffic {row['28D Transactions % Chg']:+.1f}%, avg ticket {row['28D Avg Ticket % Chg']:+.1f}%)"
            lines.append(f"  {position}. {row['Store Code']} {row['28D Net Sales % Chg']:+6.1f}%  {names.get(row['Store Code'], '')}{why}")

    ranked = analysis.ranked()
    section("TOP 3 (28-Day Net Sales % Chg)")
    ranking(ranked.head(3))
    section("BOTTOM 3 (worst first)")
    ranking(ranked.tail(3).iloc[::-1])

    section("DATA QUALITY")
    lines.append(
        f"  This run: {scan.files_seen} files examined, {len(stored.added)} new store-weeks loaded, "
        f"{len(stored.revised)} revised, {len(scan.skipped)} skipped."
    )
    for code, week in stored.revised:
        lines.append(f"  Revised: {code} {week.date()} -- a resubmitted report replaced the stored figures.")
    skipped = {category: [path for path, c, _ in scan.skipped if c == category] for category in (UNMAPPED_STORE, UNREADABLE, NOT_A_REPORT, SUPERSEDED)}
    unmapped = Counter(store_in_filename(path) for path in skipped[UNMAPPED_STORE])
    for name, count in sorted(unmapped.items()):
        lines.append(f"  Unmapped store skipped: {name} ({count} file{'s' * (count > 1)}) -- not in Store_Directory.csv.")
    for path in skipped[UNREADABLE]:
        lines.append(f"  Unreadable file skipped: {path}")
    for path in skipped[SUPERSEDED]:
        lines.append(f"  Duplicate in this drop, older copy ignored: {path}")
    if skipped[NOT_A_REPORT]:
        lines.append(f"  Not reports, ignored: {', '.join(path.rsplit('/', 1)[-1] for path in skipped[NOT_A_REPORT])}")
    for _, row in analysis.flagged.iterrows():
        lines.append(f"  Excluded impossible values, {row['Store Code']} {row['Week Ending'].date()}: {row['Data Quality Flag']}")
    for code, weeks in analysis.missing_weeks.items():
        lines.append(f"  Missing weeks (no readable report): {code} {', '.join(map(str, weeks))}")
    for _, row in analysis.stale.iterrows():
        lines.append(f"  No report this week: {row['Store Code']} -- figures above are from {row['Week Ending'].date()}.")
    lines.extend([rule, ""])
    return "\n".join(lines)


def store_in_filename(path: str) -> str:
    """The store a skipped file was about, for grouping (name or code from the filename)."""
    name = path.rsplit("/", 1)[-1]
    if match := NAME_FILE.match(name):
        return match["name"]
    match = CODE_FILE.match(name) or EXPORT_FILE.match(name)
    return match["code"].upper() if match else name


# --- Command line ------------------------------------------------------------


def log_summary(result: ScanResult, frame: pd.DataFrame) -> None:
    log.info("-" * 60)
    log.info("Files examined: %d | rows consolidated: %d | files skipped: %d", result.files_seen, len(frame), len(result.skipped))
    if not frame.empty:
        weeks = frame["Week Ending"]
        log.info("Stores: %s", ", ".join(f"{code} ({n})" for code, n in frame["Store Code"].value_counts().sort_index().items()))
        log.info("Weeks: %s to %s", weeks.min().date(), weeks.max().date())
    for category in (UNMAPPED_STORE, UNREADABLE, NOT_A_REPORT, SUPERSEDED):
        paths = [path for path, skipped_as, _ in result.skipped if skipped_as == category]
        if paths:
            log.info("Skipped, %s (%d): %s", category, len(paths), "; ".join(paths))
    flagged = frame[frame["Data Quality Flag"] != ""]
    log.info("Rows with flagged values: %d", len(flagged))
    for row in flagged.itertuples(index=False):
        log.info("  %s %s -> %s", row[0], row[1].date(), row[-1])


def configure_logging(log_file: Path | None) -> None:
    log.setLevel(logging.DEBUG)
    console = logging.StreamHandler()
    console.setLevel(logging.INFO)
    console.setFormatter(logging.Formatter("%(levelname)-7s %(message)s"))
    log.addHandler(console)
    if log_file is not None:
        # The file keeps the per-file detail (and history across weekly runs).
        to_file = logging.FileHandler(log_file, encoding="utf-8")
        to_file.setFormatter(logging.Formatter("%(asctime)s %(levelname)-7s %(message)s"))
        log.addHandler(to_file)


def main(argv: list[str] | None = None) -> int:
    here = Path(__file__).resolve().parent
    parser = argparse.ArgumentParser(description="Consolidate a weekly drop into Sales_Master.xlsx and write the ops brief.")
    parser.add_argument("folder", type=Path, help="folder to scan (subfolders and .zip files included)")
    parser.add_argument("--directory", type=Path, default=here / "Store_Directory.csv", help="store directory CSV")
    parser.add_argument("--master", type=Path, default=here / "Sales_Master.xlsx", help="master workbook (created if missing)")
    parser.add_argument("--brief", type=Path, default=here / "weekly_ops_brief.txt", help="where to write the brief")
    parser.add_argument("--log-file", type=Path, default=here / "process_reports.log", help="run log (appended to)")
    args = parser.parse_args(argv)

    configure_logging(args.log_file)
    if not args.folder.is_dir():
        log.error("%s is not a folder", args.folder)
        return 2
    if not args.directory.is_file():
        log.error("Store directory %s not found", args.directory)
        return 2

    # Part 1
    log.info("Scanning %s", args.folder.resolve())
    directory = load_directory(args.directory)
    scan = process_folder(args.folder, directory)
    new_rows = scan.to_frame()
    log_summary(scan, new_rows)

    # Part 2. If the existing master cannot be read we stop rather than
    # overwrite it -- losing the history is worse than missing one week.
    try:
        master = load_master(args.master)
    except PermissionError:
        log.error("Cannot open %s -- is it open in Excel? Close it and re-run.", args.master)
        return 1
    except Exception as exc:
        log.error("Cannot read %s (%s); not touching it. Fix or restore it and re-run.", args.master, exc)
        return 1
    stored = upsert(master, new_rows)
    try:
        save_master(stored.master, args.master)
    except PermissionError:
        log.error("Cannot write %s -- is it open in Excel? Close it and re-run.", args.master)
        return 1
    log.info(
        "Master %s: %d rows (%d added, %d revised, %d unchanged)",
        args.master.name, len(stored.master), len(stored.added), len(stored.revised), stored.unchanged,
    )
    for code, week in stored.revised:
        log.info("Revised %s %s: a resubmitted report replaced the stored row", code, week.date())

    # Parts 3 and 4
    brief = render_brief(analyse(stored.master), directory, scan, stored)
    args.brief.write_text(brief, encoding="utf-8")
    log.info("Wrote %s", args.brief.name)
    print("\n" + brief)
    return 0


if __name__ == "__main__":
    sys.exit(main())
